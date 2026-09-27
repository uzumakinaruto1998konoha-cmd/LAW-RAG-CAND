from __future__ import annotations

import sys
import unittest
from contextlib import nullcontext
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from law_rag.ingestion.errors import IngestionError, ReviewIncompleteError, ReviewItemNotFoundError
from law_rag.ingestion.extraction import ExtractionResult, SourceSpan
from law_rag.ingestion.legal_models import (
    LegalParseResult,
    MetadataField,
    NodeKind,
    RelationType,
    VerificationStatus,
)
from law_rag.ingestion.legal_parsing import parse_legal_document
from law_rag.ingestion.legal_postgres import PostgresLegalParseRepository
from law_rag.ingestion.legal_rules import LegalParseConfig, default_legal_parse_config, parse_vietnamese_date, roman_to_int
from law_rag.ingestion.legal_review import (
    ApprovedLegalDocument,
    LegalReviewService,
    ReviewAction,
    ReviewDecision,
    ReviewItemKind,
    apply_decisions,
    pending_tasks,
    require_approval,
)

NOW = datetime(2026, 9, 27, 8, 0, tzinfo=timezone.utc)

DOCUMENT_LINES: tuple[str, ...] = (
    "QUYẾT ĐỊNH",
    "Về việc phê duyệt kế hoạch",
    "Số hiệu: 184/2024/QĐ-TTg",
    "Ngày ban hành: 22/12/2024",
    "Cơ quan ban hành: Bộ Tư pháp",
    "Chương I",
    "Quy trình thực hiện",
    "Điều 1. Phạm vi áp dụng",
    "1. Điều này quy định phạm vi áp dụng.",
    "2. Căn cứ quy định tại Điều 5 của Nghị định số 12/2024/NĐ-CP, thực hiện như sau.",
    "Điều 2. Hiệu lực thi hành",
    "Quyết định này có hiệu lực từ 01/01/2025.",
)

POINT_LINES: tuple[str, ...] = (
    "Điều 1. Đối tượng áp dụng",
    "Khoản 1. Các cơ quan thực hiện",
    "a) Cơ quan thực hiện nhiệm vụ.",
    "b) Tổ chức được giao nhiệm vụ.",
)


def flow_spans(lines: tuple[str, ...], *, page: int = 1, confidence: float | None = None) -> ExtractionResult:
    spans = tuple(
        SourceSpan(page, f"page:{page}/paragraph:{index}", text, None, "docx_text", confidence)
        for index, text in enumerate(lines, start=1)
    )
    return ExtractionResult(spans, None, "ingestion-v1", (), (), False, ())


def layout_spans() -> ExtractionResult:
    spans = (
        SourceSpan(1, "page:1/span:1", "Điều 1", (10.0, 10.0, 60.0, 20.0), "pdf_text", None),
        SourceSpan(1, "page:1/span:2", ".", (60.0, 10.0, 63.0, 20.0), "pdf_text", None),
        SourceSpan(1, "page:1/span:3", "Nội dung", (70.0, 10.0, 130.0, 20.0), "pdf_text", None),
        SourceSpan(1, "page:1/span:4", "1. Điểm a", (10.0, 30.0, 90.0, 40.0), "pdf_text", None),
    )
    return ExtractionResult(spans, 1, "ingestion-v1", (), (), False, ())


class LineAssemblyTests(unittest.TestCase):
    def test_merges_layout_spans_into_one_reading_order_line(self) -> None:
        result = parse_legal_document(layout_spans(), job_id="job-1")
        article = result.nodes_of_kind(NodeKind.ARTICLE)[0]
        self.assertEqual(article.title, "Nội dung")
        self.assertEqual(article.page_numbers, (1,))
        self.assertIn("page:1/span:3", article.source_locators)


class HierarchyTests(unittest.TestCase):
    def setUp(self) -> None:
        self.result = parse_legal_document(flow_spans(DOCUMENT_LINES), job_id="job-1", source="van-ban/2024")

    def test_builds_chapter_article_clause_point_tree(self) -> None:
        kinds = [(node.kind, node.ordinal) for node in self.result.nodes]
        self.assertIn((NodeKind.CHAPTER, "1"), kinds)
        self.assertIn((NodeKind.ARTICLE, "1"), kinds)
        self.assertIn((NodeKind.CLAUSE, "1"), kinds)
        self.assertIn((NodeKind.CLAUSE, "2"), kinds)
        article_one = self.result.nodes_of_kind(NodeKind.ARTICLE)[0]
        self.assertEqual(len(article_one.child_ids), 2)
        clause = self.result.node(article_one.child_ids[0])
        assert clause is not None
        self.assertEqual(clause.kind, NodeKind.CLAUSE)

    def test_labelled_points_nest_under_their_clause(self) -> None:
        result = parse_legal_document(flow_spans(POINT_LINES), job_id="job-points")
        clause = result.nodes_of_kind(NodeKind.CLAUSE)[0]
        self.assertEqual(clause.ordinal, "1")
        self.assertEqual(clause.label, "Khoản")
        points = [result.node(child) for child in clause.child_ids]
        self.assertEqual([point.kind for point in points], [NodeKind.POINT, NodeKind.POINT])
        self.assertEqual([point.ordinal for point in points], ["a", "b"])

    def test_keeps_article_content_and_roman_chapter_canonicalisation(self) -> None:
        article = self.result.nodes_of_kind(NodeKind.ARTICLE)[0]
        self.assertEqual(article.label, "Điều")
        self.assertEqual(article.title, "Phạm vi áp dụng")
        self.assertEqual(article.provenance.source_locator, "page:1/paragraph:8")
        clause = self.result.node(article.child_ids[0])
        assert clause is not None
        self.assertIn("phạm vi áp dụng", clause.content)
        self.assertEqual(self.result.nodes_of_kind(NodeKind.CHAPTER)[0].ordinal, "1")
        self.assertEqual(self.result.nodes_of_kind(NodeKind.CHAPTER)[0].title, "Quy trình thực hiện")
        self.assertTrue(self.result.review_required)

    def test_duplicate_article_and_out_of_order_ordinal_are_reported(self) -> None:
        lines = DOCUMENT_LINES + ("Điều 2. Lặp lại", "Điều 1. Số lại xuống")
        result = parse_legal_document(flow_spans(lines), job_id="job-2")
        codes = " ".join(result.warnings)
        self.assertIn("HIERARCHY_DUPLICATE_ARTICLE_2", codes)
        self.assertIn("HIERARCHY_UNEXPECTED_ORDINAL", codes)
        flagged = [node for node in result.nodes if node.is_pending]
        self.assertTrue(flagged)

    def test_document_without_structure_routes_to_review(self) -> None:
        result = parse_legal_document(flow_spans(("Một đoạn văn bản không có cấu trúc.",)), job_id="job-3")
        self.assertEqual(result.nodes, ())
        self.assertIn("HIERARCHY_NO_NODES", result.warnings)
        self.assertTrue(result.review_required)

    def test_appendix_without_number_is_flagged_for_review(self) -> None:
        lines = DOCUMENT_LINES + ("Phụ lục", "Bảng biểu mẫu")
        result = parse_legal_document(flow_spans(lines), job_id="job-appendix")
        appendix = result.nodes_of_kind(NodeKind.APPENDIX)[0]
        self.assertEqual(appendix.ordinal, "1")
        self.assertEqual(appendix.title, "Bảng biểu mẫu")
        self.assertIn("HIERARCHY_IMPLICIT_ORDINAL_APPENDIX_1", result.warnings)
        self.assertTrue(appendix.is_pending)

    def test_orphan_list_marker_is_flagged(self) -> None:
        result = parse_legal_document(flow_spans(("1. Mục danh sách không có điều cha",)), job_id="job-4")
        self.assertIn("HIERARCHY_ORPHAN_MARKER", result.warnings)
        self.assertEqual(result.nodes, ())

    def test_parser_config_hash_is_stable_and_rule_sensitive(self) -> None:
        config = default_legal_parse_config()
        first = parse_legal_document(flow_spans(DOCUMENT_LINES), job_id="job-5", config=config)
        second = parse_legal_document(flow_spans(DOCUMENT_LINES), job_id="job-5", config=config)
        self.assertEqual(first.parser_config_hash, second.parser_config_hash)
        self.assertEqual(len(first.parser_config_hash), 64)
        stricter = LegalParseConfig(
            parser_version=config.parser_version,
            heading_rules=config.heading_rules,
            metadata_rules=config.metadata_rules,
            relation_rules=config.relation_rules,
            review_confidence_threshold=0.95,
        )
        third = parse_legal_document(flow_spans(DOCUMENT_LINES), job_id="job-5", config=stricter)
        self.assertNotEqual(first.parser_config_hash, third.parser_config_hash)
        self.assertGreaterEqual(len(third.pending_items()), len(first.pending_items()))


class MetadataTests(unittest.TestCase):
    def setUp(self) -> None:
        self.result = parse_legal_document(flow_spans(DOCUMENT_LINES), job_id="job-1", source="van-ban/2024")

    def metadata_value(self, field: MetadataField) -> str | None:
        values = [item.value for item in self.result.metadata_for(field)]
        return values[0] if values else None

    def test_extracts_structural_metadata_with_provenance(self) -> None:
        self.assertEqual(self.metadata_value(MetadataField.DOCUMENT_TYPE), "QUYẾT ĐỊNH")
        self.assertEqual(self.metadata_value(MetadataField.DOCUMENT_NUMBER), "184/2024/QĐ-TTg")
        self.assertEqual(self.metadata_value(MetadataField.ISSUE_DATE), "2024-12-22")
        self.assertEqual(self.metadata_value(MetadataField.EFFECTIVE_DATE), "2025-01-01")
        self.assertEqual(self.metadata_value(MetadataField.LANGUAGE), "vie")
        self.assertEqual(self.metadata_value(MetadataField.SOURCE), "van-ban/2024")
        number = self.result.metadata_for(MetadataField.DOCUMENT_NUMBER)[0]
        self.assertEqual(number.provenance.source_locator, "page:1/paragraph:3")
        self.assertEqual(number.provenance.page_number, 1)
        self.assertIn("184/2024/QĐ-TTg", number.raw_value)
        self.assertEqual(number.provenance.detector, "document_number_label")

    def test_issuing_body_and_title_stay_low_confidence_for_review(self) -> None:
        body = self.result.metadata_for(MetadataField.ISSUING_BODY)[0]
        self.assertEqual(body.value, "Bộ Tư pháp")
        self.assertLess(body.confidence, 0.85)
        self.assertTrue(body.is_pending)
        title = self.result.metadata_for(MetadataField.TITLE)[0]
        self.assertEqual(title.value, "Về việc phê duyệt kế hoạch")
        self.assertTrue(title.is_pending)

    def test_missing_required_fields_produce_stable_warnings(self) -> None:
        result = parse_legal_document(flow_spans(("Một đoạn văn bản.",)), job_id="job-6")
        codes = set(result.warnings)
        self.assertIn("METADATA_MISSING_DOCUMENT_NUMBER", codes)
        self.assertIn("METADATA_MISSING_ISSUE_DATE", codes)
        self.assertIn("METADATA_MISSING_DOCUMENT_TYPE", codes)

    def test_ambiguous_date_is_penalised_and_flagged(self) -> None:
        result = parse_legal_document(flow_spans(("Ngày ban hành: 05/12/2024",)), job_id="job-7")
        issue = result.metadata_for(MetadataField.ISSUE_DATE)[0]
        self.assertEqual(issue.value, "2024-12-05")
        self.assertIn("METADATA_AMBIGUOUS_ISSUE_DATE", result.warnings)
        self.assertLess(issue.confidence, 0.6)

    def test_extraction_review_warning_is_inherited(self) -> None:
        extraction = ExtractionResult(
            (SourceSpan(1, "page:1/paragraph:1", "Điều 1. Nội dung", None, "ocr", 71.0),),
            1, "ingestion-v1", (), ((1, 71.0),), True, ("OCR_LOW_CONFIDENCE_PAGE_1",),
        )
        result = parse_legal_document(extraction, job_id="job-8")
        self.assertIn("OCR_LOW_CONFIDENCE_PAGE_1", result.warnings)
        self.assertTrue(result.review_required)
        self.assertLess(result.nodes[0].confidence, 0.85)


class RelationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.result = parse_legal_document(flow_spans(DOCUMENT_LINES), job_id="job-1")

    def test_relations_are_candidates_only(self) -> None:
        relation = next(item for item in self.result.relations if item.target_document_number == "12/2024/NĐ-CP")
        self.assertEqual(relation.relation_type, RelationType.REFERENCES)
        self.assertIs(relation.verification, VerificationStatus.UNVERIFIED)
        self.assertTrue(relation.is_pending)
        self.assertFalse(relation.is_resolved)
        self.assertEqual(relation.target_scope_label, "Điều 5")
        self.assertIsNotNone(relation.source_node_id)
        self.assertEqual(relation.provenance.page_number, 1)

    def test_amendment_wording_becomes_relation_candidate(self) -> None:
        lines = DOCUMENT_LINES + ("Nghị định này được sửa đổi, bổ sung tại Điều 3 của Nghị định số 7/2023/NĐ-CP.",)
        result = parse_legal_document(flow_spans(lines), job_id="job-9")
        relation = next(item for item in result.relations if item.relation_type is RelationType.AMENDS)
        self.assertEqual(relation.target_document_number, "7/2023/NĐ-CP")
        self.assertEqual(relation.target_scope_label, "Điều 3")

    def test_unresolved_relation_target_is_reported(self) -> None:
        lines = ("Điều 1. Nội dung", "Nghị định này được sửa đổi, bổ sung theo quy định hiện hành.")
        result = parse_legal_document(flow_spans(lines), job_id="job-10")
        self.assertTrue(result.relations)
        self.assertTrue(any(code.startswith("RELATION_TARGET_UNRESOLVED") for code in result.warnings))

    def test_dates_are_not_mistaken_for_document_references(self) -> None:
        self.assertEqual(parse_vietnamese_date("31/12/2024"), ("2024-12-31", False))
        self.assertEqual(parse_vietnamese_date("07/2024"), (None, True))
        self.assertEqual(roman_to_int("XIV"), 14)
        self.assertIsNone(roman_to_int("XIZ"))


class MemoryParseRepository:
    def __init__(self) -> None:
        self.results: list[LegalParseResult] = []
        self.decisions: list[ReviewDecision] = []
        self.approvals: list[object] = []

    def save_result(self, result: LegalParseResult) -> LegalParseResult:
        self.results.append(result)
        return result

    def get_latest(self, job_id: str) -> LegalParseResult | None:
        return next((item for item in reversed(self.results) if item.job_id == job_id), None)

    def save_decisions(self, job_id: str, decisions: tuple[ReviewDecision, ...] | list[ReviewDecision]) -> None:
        self.decisions.extend(decisions)

    def save_approval(self, approval: object) -> None:
        self.approvals.append(approval)


class MemoryAudit:
    def __init__(self) -> None:
        self.events: list[dict[str, str]] = []

    def record(self, **event: str) -> None:
        self.events.append(event)


class ReviewWorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.result = parse_legal_document(flow_spans(DOCUMENT_LINES), job_id="job-1")
        self.repository = MemoryParseRepository()
        self.audit = MemoryAudit()
        self.service = LegalReviewService(repository=self.repository, audit=self.audit, clock=lambda: NOW)

    def decisions_for_everything(self, result: LegalParseResult) -> list[ReviewDecision]:
        decisions = [
            ReviewDecision(ReviewItemKind.METADATA, item.assertion_id, ReviewAction.ACCEPT, "reviewer-1")
            for item in result.metadata if item.is_pending
        ]
        decisions.extend(
            ReviewDecision(ReviewItemKind.NODE, node.node_id, ReviewAction.ACCEPT, "reviewer-1")
            for node in result.nodes if node.is_pending
        )
        decisions.extend(
            ReviewDecision(
                ReviewItemKind.RELATION, relation.relation_id, ReviewAction.ACCEPT, "reviewer-1",
                target_document_number="12/2024/NĐ-CP", target_scope_label="Điều 5",
            )
            for relation in result.relations if relation.is_pending
        )
        return decisions

    def test_pending_tasks_expose_every_reviewable_item(self) -> None:
        result = parse_legal_document(flow_spans(DOCUMENT_LINES + ("Điều 2. Lặp lại số hiệu cũ",)), job_id="job-pending")
        tasks = pending_tasks(result)
        kinds = {task.item_kind for task in tasks}
        self.assertIn(ReviewItemKind.METADATA, kinds)
        self.assertIn(ReviewItemKind.NODE, kinds)
        self.assertIn(ReviewItemKind.RELATION, kinds)
        self.assertIn(ReviewItemKind.WARNING, kinds)
        self.assertTrue(all(task.item_id for task in tasks))

    def test_approval_is_blocked_until_every_item_is_resolved(self) -> None:
        with self.assertRaises(ReviewIncompleteError):
            require_approval(self.result, reviewer_id="reviewer-1", acknowledged_warnings=self.result.warnings, now=NOW)

    def test_warning_must_be_acknowledged_before_approval(self) -> None:
        flagged = parse_legal_document(flow_spans(DOCUMENT_LINES + ("Điều 2. Lặp lại số hiệu cũ",)), job_id="job-warn")
        self.assertTrue(flagged.warnings)
        reviewed = self.service.review(flagged, self.decisions_for_everything(flagged))
        self.assertEqual(reviewed.pending_items(), ())
        with self.assertRaises(ReviewIncompleteError):
            require_approval(reviewed, reviewer_id="reviewer-1", acknowledged_warnings=(), now=NOW)
        approval = require_approval(reviewed, reviewer_id="reviewer-1", acknowledged_warnings=reviewed.warnings, now=NOW)
        self.assertTrue(approval.retrievable)
        self.assertEqual(approval.approved_by, "reviewer-1")
        self.assertEqual(approval.approved_at, NOW)
        self.assertEqual(approval.blocking_warnings(), reviewed.warnings)

    def test_review_records_reviewer_audits_and_approval(self) -> None:
        reviewed = self.service.review(self.result, self.decisions_for_everything(self.result))
        self.assertTrue(self.repository.results)
        self.assertTrue(self.repository.decisions)
        approval = self.service.approve(reviewed, reviewer_id="reviewer-1", acknowledged_warnings=reviewed.warnings)
        self.assertEqual(len(self.repository.approvals), 1)
        self.assertEqual(approval.job_id, "job-1")
        actions = {event["action"] for event in self.audit.events}
        self.assertIn("review_metadata", actions)
        self.assertIn("review_relation", actions)
        self.assertIn("approve_document", actions)
        self.assertTrue(all(event["actor_id"] == "reviewer-1" for event in self.audit.events))

    def test_correction_keeps_provenance_and_records_reviewer(self) -> None:
        title = self.result.metadata_for(MetadataField.TITLE)[0]
        updated = apply_decisions(
            self.result,
            [ReviewDecision(ReviewItemKind.METADATA, title.assertion_id, ReviewAction.CORRECT, "reviewer-2", value="Tiêu đề đã sửa")],
            now=NOW,
        )
        corrected = updated.metadata_for(MetadataField.TITLE)[0]
        self.assertEqual(corrected.value, "Tiêu đề đã sửa")
        self.assertEqual(corrected.provenance, title.provenance)
        self.assertIs(corrected.verification, VerificationStatus.CORRECTED)
        self.assertEqual(corrected.reviewer_id, "reviewer-2")
        self.assertIsNone(updated.verified_metadata_value(MetadataField.DOCUMENT_NUMBER))

    def test_rejecting_a_node_removes_it_from_the_tree(self) -> None:
        article = self.result.nodes_of_kind(NodeKind.ARTICLE)[0]
        updated = apply_decisions(
            self.result,
            [ReviewDecision(ReviewItemKind.NODE, article.node_id, ReviewAction.REJECT, "reviewer-1")],
            now=NOW,
        )
        self.assertIsNone(updated.node(article.node_id))
        self.assertNotIn(article.node_id, updated.node(article.parent_id or "") .child_ids)

    def test_unknown_item_and_warning_decision_fail_safely(self) -> None:
        with self.assertRaises(ReviewItemNotFoundError):
            apply_decisions(self.result, [ReviewDecision(ReviewItemKind.NODE, "node-9999", ReviewAction.ACCEPT, "r")], now=NOW)
        with self.assertRaises(ReviewItemNotFoundError):
            apply_decisions(self.result, [ReviewDecision(ReviewItemKind.WARNING, "HIERARCHY_NO_NODES", ReviewAction.ACCEPT, "r")], now=NOW)

    def test_relation_decision_requires_a_target(self) -> None:
        with self.assertRaises(ValueError):
            ReviewDecision(ReviewItemKind.RELATION, "rel-001", ReviewAction.ACCEPT, "reviewer-1")
        with self.assertRaises(ValueError):
            ReviewDecision(ReviewItemKind.METADATA, "meta-title-001", ReviewAction.CORRECT, "reviewer-1")


class FakeCursor:
    def __init__(self, row: tuple[Any, ...] | None = None) -> None:
        self._row = row

    def fetchone(self) -> tuple[Any, ...] | None:
        return self._row

    def fetchall(self) -> list[tuple[Any, ...]]:
        return []


class FakeConnection:
    """Records SQL and returns a single run row so insert/read paths can be asserted."""

    def __init__(self, run_row: tuple[Any, ...] | None = None) -> None:
        self.queries: list[str] = []
        self.run_row = run_row

    def execute(self, query: str, params: tuple[Any, ...] = ()) -> FakeCursor:
        self.queries.append(" ".join(query.split()))
        if "FROM legal_parse_run" in query:
            return FakeCursor(self.run_row)
        return FakeCursor()


class PostgresAdapterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.connection = FakeConnection(run_row=("run-1",))
        self.repository = PostgresLegalParseRepository(lambda: nullcontext(self.connection))
        self.result = parse_legal_document(flow_spans(DOCUMENT_LINES), job_id="job-1", source="van-ban/2024")

    def test_save_result_writes_parse_run_nodes_metadata_and_relations(self) -> None:
        self.repository.save_result(self.result)
        tables: list[str] = []
        for query in self.connection.queries:
            if "INTO" not in query:
                continue
            table = query.split("INTO ")[1].split()[0]
            if table not in tables:
                tables.append(table)
        self.assertEqual(tables, [
            "legal_parse_run", "legal_node", "metadata_assertion", "document_relation_candidate",
        ])
        self.connection.run_row = None
        self.assertIsNone(self.repository.get_latest("job-1"))

    def test_save_decisions_and_approval_require_an_existing_parse_run(self) -> None:
        decision = [ReviewDecision(ReviewItemKind.NODE, "node-0001", ReviewAction.ACCEPT, "r")]
        self.connection.run_row = None
        with self.assertRaises(IngestionError):
            self.repository.save_decisions("job-1", decision)
        with self.assertRaises(IngestionError):
            self.repository.save_approval(ApprovedLegalDocument("job-1", self.result, "r", NOW, ()))
        self.connection.run_row = ("run-1",)
        self.repository.save_decisions("job-1", decision)
        self.repository.save_approval(ApprovedLegalDocument("job-1", self.result, "r", NOW, ("METADATA_MISSING_ISSUE_DATE",)))
        self.assertTrue(any("INTO review_decision" in query for query in self.connection.queries))
        self.assertTrue(any("INTO document_approval" in query for query in self.connection.queries))


class DocxEndToEndTests(unittest.TestCase):
    def test_parses_an_extracted_docx_with_page_provenance(self) -> None:
        import tempfile

        from docx import Document as OpenDocument

        from law_rag.ingestion.extraction import extract_document

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sample.docx"
            document = OpenDocument()
            for line in DOCUMENT_LINES:
                document.add_paragraph(line)
            document.save(path)
            result = parse_legal_document(extract_document(path, original_filename="sample.docx"), job_id="job-docx")
        self.assertIsNone(result.page_count)
        articles = result.nodes_of_kind(NodeKind.ARTICLE)
        self.assertEqual([node.ordinal for node in articles], ["1", "2"])
        self.assertIsNone(articles[0].provenance.page_number)
        self.assertTrue(articles[0].provenance.source_locator.startswith("paragraph:"))
        self.assertEqual(articles[0].title, "Phạm vi áp dụng")
        self.assertIn("12/2024/NĐ-CP", " ".join(item.target_reference for item in result.relations))
        self.assertTrue(result.review_required)


if __name__ == "__main__":
    unittest.main()
