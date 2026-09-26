"""Page-aware text extraction for supported ingestion v1 file formats."""

from __future__ import annotations

import io
import logging
import math
import docx
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Literal, Protocol

import pymupdf
from docx import Document as open_docx
from docx.table import Table
from docx.text.paragraph import Paragraph
import PIL
from PIL import Image, ImageOps

from .errors import EncryptedDocumentError, ExtractionError, UploadValidationError
from .ocr import OCRProcessor, OCRResult
from .validation import MAX_DOCX_ENTRIES, MAX_DOCX_UNCOMPRESSED_BYTES, MAX_IMAGE_PIXELS

LOGGER = logging.getLogger(__name__)
MAX_PDF_PAGES = 500
PIPELINE_VERSION = "ingestion-v1"
OCR_PAGE_THRESHOLD = 90.0


@dataclass(frozen=True, slots=True)
class SourceSpan:
    page_number: int | None
    source_locator: str
    text: str
    bbox: tuple[float, float, float, float] | None
    extraction_method: Literal["pdf_text", "docx_text", "plain_text", "ocr"]
    confidence: float | None = None


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    spans: tuple[SourceSpan, ...]
    page_count: int | None
    pipeline_version: str
    tool_versions: tuple[tuple[str, str], ...]
    ocr_confidence_by_page: tuple[tuple[int, float], ...]
    review_required: bool
    warnings: tuple[str, ...]

    @property
    def text(self) -> str:
        return "\n".join(span.text for span in self.spans if span.text)


class PageRenderer(Protocol):
    def render_page(self, page: pymupdf.Page, *, max_pixels: int) -> tuple[bytes, int, int]: ...


class DefaultPageRenderer:
    """Render PDF pages at 300 DPI or lower to stay within the pixel cap."""

    def render_page(self, page: pymupdf.Page, *, max_pixels: int = MAX_IMAGE_PIXELS) -> tuple[bytes, int, int]:
        rect = page.rect
        point_area = max(1.0, rect.width * rect.height)
        scale = min(300.0 / 72.0, math.sqrt(max_pixels / point_area))
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), colorspace=pymupdf.csRGB, alpha=False)
        if pixmap.width * pixmap.height > max_pixels:
            raise ExtractionError("Rendered PDF page exceeds the safe pixel limit")
        return pixmap.tobytes("png"), pixmap.width, pixmap.height


def _ocr_spans(
    result: OCRResult,
    *,
    page_number: int,
    page_rect: tuple[float, float, float, float],
    image_size: tuple[int, int],
) -> tuple[SourceSpan, ...]:
    page_width = page_rect[2] - page_rect[0]
    page_height = page_rect[3] - page_rect[1]
    width, height = image_size
    scale_x = page_width / width if width else 1.0
    scale_y = page_height / height if height else 1.0
    return tuple(
        SourceSpan(
            page_number=page_number,
            source_locator=f"page:{page_number}/ocr-word:{index}",
            text=word.text,
            bbox=(
                page_rect[0] + word.bbox[0] * scale_x,
                page_rect[1] + word.bbox[1] * scale_y,
                page_rect[0] + word.bbox[2] * scale_x,
                page_rect[1] + word.bbox[3] * scale_y,
            ),
            extraction_method="ocr",
            confidence=word.confidence,
        )
        for index, word in enumerate(result.words, start=1)
        if word.text.strip()
    )


def _needs_ocr(text: str) -> bool:
    useful_chars = sum(char.isalnum() for char in text)
    return useful_chars < 20


def _extract_pdf(path: Path, ocr: OCRProcessor | None, renderer: PageRenderer) -> ExtractionResult:
    spans: list[SourceSpan] = []
    confidence_by_page: list[tuple[int, float]] = []
    warnings: list[str] = []
    review = False
    tool_versions: dict[str, str] = {"pymupdf": pymupdf.VersionBind}
    try:
        document = pymupdf.open(path)
    except Exception as exc:
        LOGGER.info("PDF open failed path=%s", path.name)
        raise ExtractionError("PDF file is damaged or unsupported") from exc
    with document:
        if document.needs_pass:
            raise EncryptedDocumentError("Encrypted PDF is not supported")
        page_count = document.page_count
        if page_count < 1 or page_count > MAX_PDF_PAGES:
            raise UploadValidationError(f"PDF must contain between 1 and {MAX_PDF_PAGES} pages")
        try:
            for page_index in range(page_count):
                page = document.load_page(page_index)
                page_number = page_index + 1
                page_rect = (float(page.rect.x0), float(page.rect.y0), float(page.rect.x1), float(page.rect.y1))
                page_spans: list[SourceSpan] = []
                blocks = page.get_text("dict", sort=True).get("blocks", [])
                for block_index, block in enumerate(blocks, start=1):
                    for line_index, line in enumerate(block.get("lines", []), start=1):
                        for span_index, span in enumerate(line.get("spans", []), start=1):
                            text = str(span.get("text", ""))
                            if not text.strip():
                                continue
                            box = tuple(float(value) for value in span.get("bbox", (0, 0, 0, 0)))
                            page_spans.append(SourceSpan(
                                page_number, f"page:{page_number}/block:{block_index}/line:{line_index}/span:{span_index}",
                                text, box, "pdf_text", None,
                            ))
                page_text = "".join(span.text for span in page_spans)
                if not _needs_ocr(page_text):
                    spans.extend(page_spans)
                    continue
                if ocr is None:
                    review = True
                    warnings.append(f"OCR_UNAVAILABLE_PAGE_{page_number}")
                    spans.extend(page_spans)
                    continue
                try:
                    image_bytes, pixel_width, pixel_height = renderer.render_page(page, max_pixels=MAX_IMAGE_PIXELS)
                    ocr_result = ocr.recognize(image_bytes)
                    spans.extend(_ocr_spans(ocr_result, page_number=page_number, page_rect=page_rect, image_size=(pixel_width, pixel_height)))
                    confidence_by_page.append((page_number, ocr_result.mean_confidence))
                    tool_versions.update(ocr_result.tool_versions)
                    if not ocr_result.words or ocr_result.mean_confidence < OCR_PAGE_THRESHOLD:
                        review = True
                        warnings.append(f"OCR_LOW_CONFIDENCE_PAGE_{page_number}")
                except Exception as exc:
                    review = True
                    code = exc.code if isinstance(exc, ExtractionError) else "OCR_PROCESSING_FAILURE"
                    warnings.append(f"{code}_PAGE_{page_number}")
                    spans.extend(page_spans)
        except (UploadValidationError, EncryptedDocumentError, ExtractionError):
            raise
        except Exception as exc:
            LOGGER.info("PDF page extraction failed path=%s", path.name)
            raise ExtractionError("PDF page extraction failed") from exc
    if not spans:
        review = True
        warnings.append("NO_TEXT_EXTRACTED")
    return ExtractionResult(tuple(spans), page_count, PIPELINE_VERSION, tuple(sorted(tool_versions.items())), tuple(confidence_by_page), review, tuple(warnings))


def _paragraph_break_count(paragraph: Paragraph) -> int:
    rendered = paragraph._p.xpath(".//w:lastRenderedPageBreak")
    explicit = paragraph._p.xpath(".//w:br[@w:type='page']")
    return max(len(rendered), len(explicit))


def _extract_docx(path: Path) -> ExtractionResult:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES or sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise UploadValidationError("DOCX package exceeds safe expansion limits")
            if archive.testzip() is not None:
                raise ExtractionError("DOCX package integrity check failed")
        document = open_docx(path)
    except (UploadValidationError, ExtractionError):
        raise
    except Exception as exc:
        raise ExtractionError("DOCX file is damaged or unsupported") from exc

    spans: list[SourceSpan] = []
    page_number = 1
    paragraph_number = 0
    table_number = 0
    has_page_break = False

    def add_paragraph(paragraph: Paragraph, locator: str) -> None:
        nonlocal page_number, has_page_break
        text = paragraph.text
        if text:
            spans.append(SourceSpan(page_number, locator, text, None, "docx_text", None))
        breaks = _paragraph_break_count(paragraph)
        if breaks:
            has_page_break = True
            page_number += breaks

    def add_table(table: Table, locator: str) -> None:
        nonlocal page_number
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                for paragraph_index, paragraph in enumerate(cell.paragraphs, start=1):
                    add_paragraph(paragraph, f"{locator}/row:{row_index}/cell:{cell_index}/paragraph:{paragraph_index}")
                for nested_index, nested in enumerate(cell.tables, start=1):
                    add_table(nested, f"{locator}/row:{row_index}/cell:{cell_index}/table:{nested_index}")

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_number += 1
            add_paragraph(block, f"paragraph:{paragraph_number}")
        elif isinstance(block, Table):
            table_number += 1
            add_table(block, f"table:{table_number}")
    if not has_page_break:
        spans = [replace(span, page_number=None) for span in spans]
    return ExtractionResult(tuple(spans), None, PIPELINE_VERSION, (("python-docx", docx.__version__),), (), not bool(spans), ("NO_TEXT_EXTRACTED",) if not spans else ())


def _extract_text(path: Path) -> ExtractionResult:
    try:
        text = path.read_text(encoding="utf-8-sig", errors="strict")
    except UnicodeDecodeError as exc:
        raise ExtractionError("Text file must be valid UTF-8") from exc
    except OSError as exc:
        raise ExtractionError("Text file could not be read") from exc
    pages = text.split("\f")
    spans = tuple(SourceSpan(index, f"page:{index}", content, None, "plain_text", None) for index, content in enumerate(pages, start=1) if content)
    return ExtractionResult(spans, len(pages), PIPELINE_VERSION, (), (), not bool(spans), ("NO_TEXT_EXTRACTED",) if not spans else ())


def _extract_image(path: Path, ocr: OCRProcessor | None) -> ExtractionResult:
    try:
        with Image.open(path) as image:
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise UploadValidationError("Image exceeds the 20 megapixel limit")
            image.load()
            rgb = ImageOps.exif_transpose(image).convert("RGB")
            width, height = rgb.size
            buffer = io.BytesIO()
            rgb.save(buffer, format="PNG")
            image_bytes = buffer.getvalue()
    except UploadValidationError:
        raise
    except Exception as exc:
        raise ExtractionError("Image file is damaged or unsupported") from exc
    if ocr is None:
        return ExtractionResult((), 1, PIPELINE_VERSION, (("Pillow", PIL.__version__),), (), True, ("OCR_UNAVAILABLE_PAGE_1",))
    try:
        result = ocr.recognize(image_bytes)
        spans = _ocr_spans(result, page_number=1, page_rect=(0.0, 0.0, float(width), float(height)), image_size=(width, height))
        review = not spans or result.mean_confidence < OCR_PAGE_THRESHOLD
        warnings = ("OCR_LOW_CONFIDENCE_PAGE_1",) if review else ()
        versions = (("Pillow", PIL.__version__),) + result.tool_versions
        return ExtractionResult(spans, 1, PIPELINE_VERSION, tuple(sorted(dict(versions).items())), ((1, result.mean_confidence),), review, warnings)
    except ExtractionError as exc:
        return ExtractionResult((), 1, PIPELINE_VERSION, (("Pillow", PIL.__version__),), (), True, (exc.code,))


def extract_document(
    path: Path,
    *,
    original_filename: str,
    ocr: OCRProcessor | None = None,
    renderer: PageRenderer | None = None,
) -> ExtractionResult:
    """Extract text and provenance without modifying the immutable source blob."""
    extension = Path(original_filename.replace("\\", "/")).suffix.lower()
    if extension == ".pdf":
        result = _extract_pdf(path, ocr, renderer or DefaultPageRenderer())
    elif extension == ".docx":
        result = _extract_docx(path)
    elif extension == ".txt":
        result = _extract_text(path)
    elif extension in {".jpg", ".jpeg", ".png"}:
        result = _extract_image(path, ocr)
    else:
        raise ExtractionError("File format is not supported for extraction")
    LOGGER.info(
        "Document extracted extension=%s pages=%s spans=%d review_required=%s pipeline=%s",
        extension, result.page_count, len(result.spans), result.review_required, result.pipeline_version,
    )
    return result
