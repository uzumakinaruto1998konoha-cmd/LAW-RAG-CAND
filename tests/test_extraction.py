from __future__ import annotations

import io
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import pymupdf
from docx import Document
from PIL import Image

from law_rag.ingestion.errors import EncryptedDocumentError, ExtractionError, OCRUnavailableError, UploadValidationError
from law_rag.ingestion.extraction import extract_document
from law_rag.ingestion.ocr import OCRResult, OCRWord, TesseractOCR


class FakeOCR:
    def __init__(self, confidence: float = 95.0, text: str = "recognized words") -> None:
        self.confidence = confidence
        self.text = text
        self.calls = 0

    def recognize(self, image_bytes: bytes) -> OCRResult:
        self.calls += 1
        words = (OCRWord(self.text, self.confidence, (10.0, 20.0, 80.0, 40.0)),) if self.text else ()
        return OCRResult(words, self.confidence if words else 0.0, (("tesseract", "tesseract 5.test"), ("traineddata", "vie:sha256:test")))


class ExtractionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def make_pdf(self, name: str, *, pages: tuple[str, ...], encrypted: bool = False) -> Path:
        document = pymupdf.open()
        for text in pages:
            page = document.new_page(width=300, height=400)
            if text:
                page.insert_text((30, 50), text)
        path = self.root / name
        if encrypted:
            document.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
        else:
            document.save(path)
        document.close()
        return path

    def test_extracts_pdf_text_with_page_and_bbox(self) -> None:
        path = self.make_pdf("text.pdf", pages=("A sufficiently long source page for extraction.",))
        result = extract_document(path, original_filename="text.pdf")
        self.assertEqual(result.page_count, 1)
        self.assertIn("sufficiently long", result.text)
        self.assertEqual(result.spans[0].page_number, 1)
        self.assertIsNotNone(result.spans[0].bbox)
        self.assertFalse(result.review_required)

    def test_pdf_ocr_runs_only_on_scan_pages_and_maps_confidence(self) -> None:
        path = self.make_pdf("mixed.pdf", pages=("This is a page with a searchable text layer.", ""))
        ocr = FakeOCR(confidence=87.0)
        result = extract_document(path, original_filename="mixed.pdf", ocr=ocr)
        self.assertEqual(ocr.calls, 1)
        self.assertIn((2, 87.0), result.ocr_confidence_by_page)
        self.assertTrue(result.review_required)
        self.assertTrue(any("LOW_CONFIDENCE_PAGE_2" in warning for warning in result.warnings))
        ocr_span = next(span for span in result.spans if span.extraction_method == "ocr")
        self.assertEqual(ocr_span.page_number, 2)
        self.assertIsNotNone(ocr_span.bbox)

    def test_empty_pdf_page_without_ocr_routes_to_review(self) -> None:
        path = self.make_pdf("scan.pdf", pages=("",))
        result = extract_document(path, original_filename="scan.pdf")
        self.assertTrue(result.review_required)
        self.assertIn("OCR_UNAVAILABLE_PAGE_1", result.warnings)

    def test_rejects_encrypted_pdf(self) -> None:
        path = self.make_pdf("locked.pdf", pages=("text",), encrypted=True)
        with self.assertRaises(EncryptedDocumentError):
            extract_document(path, original_filename="locked.pdf")

    def test_rejects_pdf_above_page_limit(self) -> None:
        path = self.make_pdf("many.pdf", pages=tuple("" for _ in range(501)))
        with self.assertRaises(UploadValidationError):
            extract_document(path, original_filename="many.pdf")

    def test_extracts_docx_paragraphs_and_tables_in_order(self) -> None:
        document = Document()
        document.add_paragraph("first paragraph")
        document.add_page_break()
        document.add_paragraph("second paragraph")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "table cell text"
        path = self.root / "sample.docx"
        document.save(path)
        result = extract_document(path, original_filename="sample.docx")
        self.assertEqual(result.text.splitlines(), ["first paragraph", "second paragraph", "table cell text"])
        self.assertEqual(result.spans[0].page_number, 1)
        self.assertEqual(result.spans[1].page_number, 2)
        self.assertTrue(result.spans[2].source_locator.startswith("table:1/"))

    def test_docx_without_page_breaks_uses_source_locator_without_fabricated_page(self) -> None:
        document = Document()
        document.add_paragraph("page not available")
        path = self.root / "no-layout.docx"
        document.save(path)
        result = extract_document(path, original_filename="no-layout.docx")
        self.assertIsNone(result.spans[0].page_number)
        self.assertEqual(result.spans[0].source_locator, "paragraph:1")

    def test_extracts_utf8_text_bom_and_form_feed_pages(self) -> None:
        path = self.root / "sample.txt"
        path.write_text("\ufefffirst page\fsecond page", encoding="utf-8")
        result = extract_document(path, original_filename="sample.txt")
        self.assertEqual(result.page_count, 2)
        self.assertEqual([span.page_number for span in result.spans], [1, 2])
        self.assertEqual(result.text, "first page\nsecond page")

    def test_extracts_image_with_ocr_and_review_threshold(self) -> None:
        path = self.root / "sample.png"
        Image.new("RGB", (200, 100), "white").save(path)
        ocr = FakeOCR(confidence=91.0, text="image text")
        result = extract_document(path, original_filename="sample.png", ocr=ocr)
        self.assertEqual(result.spans[0].page_number, 1)
        self.assertEqual(result.spans[0].bbox, (10.0, 20.0, 80.0, 40.0))
        self.assertFalse(result.review_required)
        self.assertEqual(result.ocr_confidence_by_page, ((1, 91.0),))

    def test_missing_local_tesseract_returns_stable_error(self) -> None:
        ocr = TesseractOCR(executable="law-rag-test-binary-does-not-exist")
        with self.assertRaises(OCRUnavailableError):
            ocr.recognize(b"not an image")

    def test_tesseract_tsv_preserves_word_confidence_and_box(self) -> None:
        executable = self.root / "tesseract.exe"
        tessdata = self.root / "tessdata"
        tessdata.mkdir()
        model = tessdata / "vie.traineddata"
        model.write_bytes(b"mock traineddata")
        tsv = "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n5\t1\t1\t1\t1\t1\t2\t3\t20\t8\t93.5\txinchao\n"
        with patch("law_rag.ingestion.ocr.shutil.which", return_value=str(executable)), patch(
            "law_rag.ingestion.ocr.subprocess.run",
            side_effect=[
                SimpleNamespace(returncode=0, stdout="tesseract 5.5.3\n"),
                SimpleNamespace(returncode=0, stdout="List of available languages (1):\nvie\n"),
                SimpleNamespace(returncode=0, stdout=tsv),
            ],
        ):
            result = TesseractOCR().recognize(self.image_bytes())
        self.assertEqual(result.words[0].text, "xinchao")
        self.assertEqual(result.words[0].confidence, 93.5)
        self.assertEqual(result.words[0].bbox, (2.0, 3.0, 22.0, 11.0))
        self.assertIn(hashlib.sha256(b"mock traineddata").hexdigest(), result.tool_versions[1][1])

    def image_bytes(self) -> bytes:
        image_buffer = io.BytesIO()
        Image.new("RGB", (4, 4), "white").save(image_buffer, format="PNG")
        return image_buffer.getvalue()

    def test_unsupported_extension_fails_safely(self) -> None:
        path = self.root / "sample.bin"
        path.write_bytes(b"data")
        with self.assertRaises(ExtractionError):
            extract_document(path, original_filename="sample.bin")


if __name__ == "__main__":
    unittest.main()
